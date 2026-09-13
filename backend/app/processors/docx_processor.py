from io import BytesIO
from docx import Document
from .base_processor import BaseProcessor, ParsedDocument

class DOCXProcessor(BaseProcessor):
    def process(self, data, filename):
        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
        return ParsedDocument(text, tables=tables)
