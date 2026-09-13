from io import BytesIO
from openpyxl import load_workbook
from .base_processor import BaseProcessor, ParsedDocument

class XLSXProcessor(BaseProcessor):
    def process(self, data, filename):
        wb = load_workbook(BytesIO(data), data_only=False)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
        return ParsedDocument("\n".join(parts), metadata={"sheets": wb.sheetnames})
