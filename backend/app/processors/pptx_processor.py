from io import BytesIO
from pptx import Presentation
from .base_processor import BaseProcessor, ParsedDocument

class PPTXProcessor(BaseProcessor):
    def process(self, data, filename):
        prs = Presentation(BytesIO(data))
        pages = []
        for i, slide in enumerate(prs.slides):
            text = "\n".join(s.text for s in slide.shapes if hasattr(s, "text"))
            pages.append({"slide": i + 1, "text": text})
        return ParsedDocument("\n".join(p["text"] for p in pages), pages=pages)
